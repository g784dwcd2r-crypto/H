"""An error report over the arithmetic checks: which rule fails, how badly, and for whom.

`check-tolerance` answers "is the tolerance hiding breaks". This answers the next question: of the
checks that fail, which rule (balance sheet balances, gross profit, EPS, ...), how far past the line,
how many distinct companies, and the worst offenders by name. It is the report to skim when the
headline says ~9 % of checks fail and you want to know where.

Read-only. Run it against a local lake (a remote lake scans whole tables; see steps.md, Part 9).
"""

from __future__ import annotations

from collections import defaultdict
from typing import Any

from filings_hub.ingest import checks as chk
from filings_hub.ingest.checks import EPS_APPROX_RELATIVE_TOLERANCE, EPS_RELATIVE_TOLERANCE, RELATIVE_TOLERANCE
from filings_hub.lake import layout
from filings_hub.lake.duck import Duck
from filings_hub.lake.storage import Storage

# A failing check, by how far past its tolerance line it sits. q = (relative gap) / (that check's
# tolerance), so q >= 1 means it failed; the bands say by how much.
FAIL_BANDS = ("just-over", "2-10x", ">10x")

# Plain-language names for the checks, so the report reads without knowing the concept codes.
CHECK_LABELS = {
    "assets_eq_liabilities_and_equity": "Balance sheet balances (Assets = Liabilities + Equity)",
    "assets_eq_liabilities_plus_equity": "Balance sheet balances (older layout)",
    "gross_profit": "Gross profit (Revenue - Cost = Gross profit)",
    "income_after_tax": "Income after tax (Pretax - Tax = Continuing income)",
    # kept under 52 characters, the width of the report's label column
    "eps_basic": "EPS basic (company's own numerator)",
    "eps_diluted": "EPS diluted (company's own numerator)",
    "eps_basic_approx": "EPS basic, approximate (inferred numerator, 5 %)",
    "eps_diluted_approx": "EPS diluted, approximate (inferred numerator, 5 %)",
    "net_income_is_equals_cf": "Net income agrees: income statement = cash flow",
    "ending_cash_cf_equals_bs": "Ending cash agrees: cash flow = balance sheet",
}


def _tol_expr() -> str:
    return (
        f"CASE WHEN check_name LIKE '%\\_approx' ESCAPE '\\' THEN {EPS_APPROX_RELATIVE_TOLERANCE} "
        f"WHEN check_name LIKE 'eps\\_%' ESCAPE '\\' THEN {EPS_RELATIVE_TOLERANCE} "
        f"ELSE {RELATIVE_TOLERANCE} END"
    )


# Why a check failed, read off the two sides alone. Not a diagnosis of the filing, a sorting of the
# pile: a sign flip or a factor of a thousand is a filer's tagging anomaly to list, not a check to
# fix; "just over" is a tolerance question; "unexplained" is where the real investigation goes.
REASONS = (
    "one side is zero",
    "sign: the two sides are exact negatives",
    "scale: off by a factor of 1,000 or 1,000,000",
    "period: off by a factor of 2 to 4",
    "just over the tolerance",
    "unexplained",
)


def _base_sql() -> str:
    """The failing checks with their gap, tolerance multiple `q`, ratio and reason, as CTE `r`."""
    return f"""
        WITH t AS (
            SELECT check_name, cik, accession, statement, lhs, rhs, passed, detail,
                   abs(coalesce(difference, lhs - rhs)) AS gap,
                   greatest(abs(lhs), abs(rhs)) AS mag,
                   {_tol_expr()} AS tol
            FROM sc WHERE lhs IS NOT NULL AND rhs IS NOT NULL
        ),
        q AS (
            SELECT *, CASE WHEN mag = 0 OR mag IS NULL THEN NULL ELSE (gap / mag) / tol END AS q,
                   CASE WHEN lhs = 0 OR rhs = 0 THEN NULL ELSE lhs / rhs END AS ratio
            FROM t
        ),
        r AS (
            SELECT *, CASE
                WHEN passed THEN NULL
                WHEN lhs = 0 OR rhs = 0 THEN '{REASONS[0]}'
                WHEN abs(lhs + rhs) <= 0.01 * mag THEN '{REASONS[1]}'
                WHEN abs(ratio - 1000) <= 50 OR abs(ratio - 1e6) <= 5e4
                  OR abs(ratio - 0.001) <= 5e-5 OR abs(ratio - 1e-6) <= 5e-8 THEN '{REASONS[2]}'
                WHEN abs(ratio - 2) <= 0.2 OR abs(ratio - 3) <= 0.3 OR abs(ratio - 4) <= 0.4
                  OR abs(ratio - 0.5) <= 0.05 OR abs(ratio - 0.333) <= 0.03 OR abs(ratio - 0.25) <= 0.03
                  THEN '{REASONS[3]}'
                WHEN q < 2 THEN '{REASONS[4]}'
                ELSE '{REASONS[5]}'
            END AS reason
            FROM q
        )
    """


def failure_report(storage: Storage, examples: int = 20) -> dict[str, Any]:
    duck = Duck(storage)
    try:
        if not duck.view("sc", f"{layout.STATEMENT_CHECKS}/*/*.parquet"):
            return {"total": 0, "message": "no statement_checks in the lake"}
        have_companies = duck.view("companies", layout.COMPANIES, hive=False)

        base = _base_sql()

        per_check = duck.fetch_dicts(
            base
            + """
            SELECT check_name,
                   count(*) AS total,
                   count(*) FILTER (WHERE passed) AS passed,
                   count(*) FILTER (WHERE NOT passed) AS failed,
                   count(*) FILTER (WHERE NOT passed AND q < 2) AS just_over,
                   count(*) FILTER (WHERE NOT passed AND q >= 2 AND q < 10) AS mid,
                   count(*) FILTER (WHERE NOT passed AND (q >= 10 OR q IS NULL)) AS far
            FROM r GROUP BY check_name ORDER BY failed DESC
            """
        )
        reasons = duck.fetch_dicts(
            base + "SELECT check_name, reason, count(*) AS n FROM r WHERE NOT passed GROUP BY 1, 2 ORDER BY 1, 3 DESC"
        )

        totals = duck.fetch_dicts(
            "SELECT count(*) AS checks, count(*) FILTER (WHERE NOT passed) AS failed, "
            "count(DISTINCT CASE WHEN NOT passed THEN cik END) AS companies_with_failure, "
            "count(DISTINCT cik) AS companies_checked "
            "FROM sc WHERE lhs IS NOT NULL AND rhs IS NOT NULL"
        )[0]

        name_col = "c.name" if have_companies else "NULL"
        name_join = "LEFT JOIN companies c ON c.cik = q.cik" if have_companies else ""
        worst = duck.fetch_dicts(
            base
            + f"""
            SELECT q.check_name, q.cik, {name_col} AS name, q.accession, q.lhs, q.rhs, q.q AS q, q.reason
            FROM r q {name_join}
            WHERE NOT q.passed AND q.q IS NOT NULL
            ORDER BY q.q DESC LIMIT {int(examples)}
            """
        )

        return {
            "total": totals["checks"],
            "failed": totals["failed"],
            "companies_checked": totals["companies_checked"],
            "companies_with_failure": totals["companies_with_failure"],
            "per_check": per_check,
            "reasons": reasons,
            "worst": worst,
        }
    finally:
        duck.close()


def export_failures(storage: Storage, path: str) -> int:
    """Every failing check, with the company's name and the reason, to a CSV: the named list. Returns
    the number of rows written."""
    duck = Duck(storage)
    try:
        if not duck.view("sc", f"{layout.STATEMENT_CHECKS}/*/*.parquet"):
            return 0
        have_companies = duck.view("companies", layout.COMPANIES, hive=False)
        name_col = "c.name" if have_companies else "NULL"
        name_join = "LEFT JOIN companies c ON c.cik = q.cik" if have_companies else ""
        # the CTEs go inside the copied query: COPY does not take a leading WITH
        duck.sql(
            f"""
            COPY (
                {_base_sql()}
                SELECT q.cik, {name_col} AS name, q.accession, q.statement, q.check_name, q.reason,
                       q.lhs, q.rhs, round(q.q, 2) AS tolerance_multiple, q.detail
                FROM r q {name_join}
                WHERE NOT q.passed
                ORDER BY q.reason, q.check_name, q.cik, q.accession
            ) TO '{path}' (FORMAT CSV, HEADER)
            """
        )
        return int(duck.fetch_value(_base_sql() + "SELECT count(*) FROM r WHERE NOT passed"))
    finally:
        duck.close()


def format_failure_report(report: dict[str, Any]) -> str:
    if report.get("total", 0) == 0:
        return report.get("message", "no checks found")
    checked = report["companies_checked"] or 1
    out = [
        f"Checks: {report['total']:,}   failed: {report['failed']:,} "
        f"({report['failed'] / max(report['total'], 1) * 100:.1f} %)",
        f"Companies checked: {report['companies_checked']:,}   "
        f"with at least one failing check: {report['companies_with_failure']:,} "
        f"({report['companies_with_failure'] / checked * 100:.1f} %)",
        "",
        "By check (worst first):",
        f"  {'check':<52} {'ran':>9} {'fail':>9} {'fail%':>6}  {'just':>7} {'2-10x':>7} {'>10x':>8}",
    ]
    by_check: dict[str, list[dict[str, Any]]] = {}
    for row in report.get("reasons", []):
        by_check.setdefault(row["check_name"], []).append(row)
    for r in report["per_check"]:
        label = CHECK_LABELS.get(r["check_name"], r["check_name"])
        rate = f"{r['failed'] / r['total'] * 100:.0f}%" if r["total"] else "-"
        out.append(
            f"  {label:<52.52} {r['total']:>9,} {r['failed']:>9,} {rate:>6}  "
            f"{r['just_over']:>7,} {r['mid']:>7,} {r['far']:>8,}"
        )
        if r["failed"] and by_check.get(r["check_name"]):
            parts = [f"{x['reason']} {x['n']:,}" for x in by_check[r["check_name"]]]
            out.append(f"      why: {' | '.join(parts)}")
    if report["worst"]:
        out.append("")
        out.append(f"Worst {len(report['worst'])} failures (most out of line):")
        for w in report["worst"]:
            who = w.get("name") or f"CIK {w['cik']}"
            tol = (
                EPS_APPROX_RELATIVE_TOLERANCE
                if w["check_name"].endswith("_approx")
                else EPS_RELATIVE_TOLERANCE
                if w["check_name"].startswith("eps")
                else RELATIVE_TOLERANCE
            )
            pct = (w["q"] or 0) * tol
            out.append(
                f"  cik {w['cik']:<8} {who[:28]:<28} {CHECK_LABELS.get(w['check_name'], w['check_name'])[:30]:<30} "
                f"{w['lhs']:>16,.0f} vs {w['rhs']:>16,.0f}  ({pct * 100:.1f}% off)  [{w.get('reason') or ''}]"
            )
    return "\n".join(out)


def explain(storage: Storage, cik: int, accessions: int = 3) -> dict[str, Any]:
    """The smoking gun for one company: its failing checks, and the actual statement lines behind
    them (raw `value`, the display-adjusted `value_presented`, and the `negating` flag), so we can
    see whether a sign inversion is in the stored data or in the check's comparison.

    Scoped to one company, so it is quick even on a remote lake.
    """
    from filings_hub.ingest.checks import CHECK_CONCEPTS

    duck = Duck(storage)
    try:
        sc = f"{layout.STATEMENT_CHECKS}/cik={int(cik)}"
        st = f"{layout.STATEMENTS}/cik={int(cik)}"
        if not duck.view("sc", f"{sc}/*.parquet"):
            return {"cik": cik, "message": "no statement_checks for this company"}
        duck.view("st", f"{st}/*.parquet")

        fails = duck.fetch_dicts(
            "SELECT accession, statement, check_name, lhs, rhs, difference, detail "
            "FROM sc WHERE NOT passed ORDER BY accession, check_name"
        )
        bad_accessions = list(dict.fromkeys(f["accession"] for f in fails))[:accessions]
        concepts = "', '".join(sorted(CHECK_CONCEPTS))
        lines = (
            duck.fetch_dicts(
                f"""
            SELECT accession, statement, concept, value, value_presented, negating, period_end
            FROM st
            WHERE is_primary_period AND coalesce(segments, '') = '' AND concept IN ('{concepts}')
              AND accession IN ('{"', '".join(bad_accessions)}')
            ORDER BY accession, statement, concept
            """
            )
            if bad_accessions
            else []
        )
        return {"cik": cik, "fails": fails, "lines": lines, "shown_accessions": bad_accessions}
    finally:
        duck.close()


def format_explain(report: dict[str, Any]) -> str:
    if "message" in report:
        return report["message"]
    out = [f"CIK {report['cik']}: {len(report['fails'])} failing checks"]
    for f in report["fails"][:30]:
        out.append(f"  {f['accession']}  {f['check_name']:<34} lhs {f['lhs']:>18,.0f}  rhs {f['rhs']:>18,.0f}")
    out.append("")
    out.append("Statement lines behind the shown filings (value = as filed, presented = display sign):")
    out.append(f"  {'concept':<48} {'stmt':>4} {'negating':>8} {'value':>18} {'presented':>18}")
    for ln in report["lines"]:
        out.append(
            f"  {ln['concept'][:48]:<48} {ln['statement']:>4} {ln['negating']!s:>8} "
            f"{(ln['value'] or 0):>18,.0f} {(ln['value_presented'] or 0):>18,.0f}"
        )
    return "\n".join(out)


# -- digging into one check ----------------------------------------------------------------------


def _quoted(values: list[str]) -> str:
    return ", ".join("'" + v.replace("'", "''") + "'" for v in values)


def _fmt(v: float | None) -> str:
    if v is None:
        return "-"
    return f"{v:,.0f}" if abs(v) >= 1000 else f"{v:.4g}"


def _gap_concepts(duck: Duck, base: str, check_name: str) -> list[dict[str, Any]]:
    """Name the gap: for every failing row, which line in the same filing is worth exactly that much.

    Every cause found so far was a line the check either left out or counted twice, and in both cases
    the shortfall equals that line (or, when counted twice, half of it). Asking the lake which concept
    the gap is equal to turns "these numbers disagree" into a name, across all the failures at once
    rather than the fifteen a sample happens to show.
    """
    return duck.fetch_dicts(
        base
        + """
        , fail AS (
            SELECT accession, lhs - rhs AS gap FROM r
            WHERE check_name = ? AND NOT passed AND lhs IS NOT NULL AND rhs IS NOT NULL AND lhs <> rhs
        )
        , vals AS (
            SELECT s.accession, s.concept, s.value FROM st s
            WHERE s.accession IN (SELECT accession FROM fail)
              AND s.is_primary_period AND NOT s.is_parenthetical AND coalesce(s.segments, '') = ''
              AND s.value IS NOT NULL AND s.value <> 0
        )
        , hits AS (
            SELECT v.concept,
                   count(*) FILTER (WHERE abs(v.value - f.gap) <= 0.01 * abs(f.gap)) AS is_gap,
                   count(*) FILTER (WHERE abs(v.value + f.gap) <= 0.01 * abs(f.gap)) AS is_minus_gap,
                   count(*) FILTER (WHERE abs(2 * v.value - f.gap) <= 0.01 * abs(f.gap)) AS is_half_gap
            FROM fail f JOIN vals v ON v.accession = f.accession
            GROUP BY 1
        )
        SELECT concept, is_gap, is_minus_gap, is_half_gap, (SELECT count(*) FROM fail) AS failures
        FROM hits WHERE is_gap + is_minus_gap + is_half_gap > 0
        ORDER BY is_gap + is_minus_gap + is_half_gap DESC LIMIT 20
        """,
        [check_name],
    )


def dig(storage: Storage, check_name: str, examples: int = 15) -> dict[str, Any]:
    """Everything about one check's failures that finds the next cause, in one pass.

    Three views, the same three that found every cause so far: which line each side of the check
    used and how often each fails (the EPS and tax-identity bugs were one line failing five times
    more than the others); the reasons; and a random sample of failing filings with every check
    concept they carry, so the arithmetic can be read rather than theorised about."""
    duck = Duck(storage)
    try:
        if not duck.view("sc", f"{layout.STATEMENT_CHECKS}/*/*.parquet"):
            return {"check": check_name, "message": "no statement_checks in the lake"}
        have_companies = duck.view("companies", layout.COMPANIES, hive=False)
        base = _base_sql()
        by_rhs = duck.fetch_dicts(
            base
            + """
            SELECT split_part(detail, ' = ', 2) AS line, count(*) AS ran, count(*) FILTER (WHERE NOT passed) AS failed
            FROM r WHERE check_name = ? GROUP BY 1 ORDER BY 2 DESC
            """,
            [check_name],
        )
        by_lhs = duck.fetch_dicts(
            base
            + """
            SELECT split_part(detail, ' = ', 1) AS line, count(*) AS ran, count(*) FILTER (WHERE NOT passed) AS failed
            FROM r WHERE check_name = ? GROUP BY 1 ORDER BY 2 DESC
            """,
            [check_name],
        )
        reasons = duck.fetch_dicts(
            base + "SELECT reason, count(*) AS n FROM r WHERE check_name = ? AND NOT passed GROUP BY 1 ORDER BY 2 DESC",
            [check_name],
        )
        have_st = duck.view("st", f"{layout.STATEMENTS}/*/*.parquet")
        gap_concepts = _gap_concepts(duck, base, check_name) if have_st else []
        name_col = "c.name" if have_companies else "NULL"
        name_join = "LEFT JOIN companies c ON c.cik = q.cik" if have_companies else ""
        sample = duck.fetch_dicts(
            base
            + f"""
            SELECT q.cik, {name_col} AS name, q.accession, q.statement, q.lhs, q.rhs, q.ratio, q.reason, q.detail
            FROM r q {name_join}
            WHERE q.check_name = ? AND NOT q.passed
            ORDER BY random() LIMIT {int(examples)}
            """,
            [check_name],
        )
        for row in sample:
            row["form"], row["values"] = None, []
        if sample and have_st:
            rows = duck.fetch_dicts(
                f"""
                SELECT accession, form, statement, concept, value FROM st
                WHERE accession IN ({_quoted([r["accession"] for r in sample])})
                  AND is_primary_period AND NOT is_parenthetical AND coalesce(segments, '') = ''
                  AND value IS NOT NULL AND concept IN ({_quoted(sorted(chk.CHECK_CONCEPTS))})
                ORDER BY accession, statement, concept
                """
            )
            values: dict[str, list[str]] = defaultdict(list)
            forms: dict[str, str] = {}
            for r in rows:
                values[r["accession"]].append(f"{r['statement']}:{r['concept']}={_fmt(r['value'])}")
                forms[r["accession"]] = r["form"]
            for row in sample:
                row["form"] = forms.get(row["accession"])
                row["values"] = values.get(row["accession"], [])
        return {
            "check": check_name,
            "by_rhs": by_rhs,
            "by_lhs": by_lhs,
            "reasons": reasons,
            "gap_concepts": gap_concepts,
            "sample": sample,
        }
    finally:
        duck.close()


def format_dig(report: dict[str, Any]) -> str:
    if "message" in report:
        return report["message"]
    label = CHECK_LABELS.get(report["check"], report["check"])
    out = [f"{label}", ""]
    for title, rows in (
        ("By the line on the right of '=' ", report["by_rhs"]),
        ("By the lines on the left of '=' ", report["by_lhs"]),
    ):
        out.append(f"{title}(ran / failed):")
        for r in rows:
            rate = f"{r['failed'] / r['ran'] * 100:.1f}%" if r["ran"] else "-"
            out.append(f"  {(r['line'] or '(none)')[:88]:<88} {r['ran']:>9,} {r['failed']:>8,} {rate:>7}")
        out.append("")
    if report["reasons"]:
        out.append("Why: " + " | ".join(f"{r['reason']} {r['n']:,}" for r in report["reasons"]))
        out.append("")
    if report.get("gap_concepts"):
        total = report["gap_concepts"][0]["failures"] or 1
        out.append("What the gap is worth (share of failures where this line equals it):")
        out.append(f"  {'line':<64} {'= gap':>8} {'= -gap':>8} {'= gap/2':>8}")
        for g in report["gap_concepts"]:
            share = (g["is_gap"] + g["is_minus_gap"] + g["is_half_gap"]) / total * 100
            out.append(
                f"  {g['concept'][:64]:<64} {g['is_gap']:>8,} {g['is_minus_gap']:>8,} "
                f"{g['is_half_gap']:>8,}  {share:>5.1f}%"
            )
        out.append("")
    if report["sample"]:
        out.append(f"{len(report['sample'])} failing filings, at random (computed vs reported):")
        for s in report["sample"]:
            who = s.get("name") or f"CIK {s['cik']}"
            ratio = f"{s['ratio']:.3f}" if s.get("ratio") is not None else "-"
            out.append(
                f"  {s['accession']}  {who[:30]:<30} {s.get('form') or '':<6} "
                f"{_fmt(s['lhs'])} vs {_fmt(s['rhs'])}  (ratio {ratio})  [{s['reason']}]"
            )
            out.append(f"      {s['detail']}")
            if s["values"]:
                out.append("      " + "  ".join(s["values"]))
    return "\n".join(out)


def export_failures_xlsx(storage: Storage, path: str) -> int:
    """Every failing check as a workbook, for a reviewer rather than a machine.

    One sheet summarising by check and by reason, one sheet of every failure with the company named,
    frozen headers and a filter over every column. CIK and accession are written as text so a
    spreadsheet does not turn them into scientific notation, which is what makes a CSV of this
    painful to read. Returns the number of failing checks written.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    report = failure_report(storage, examples=0)
    if not report.get("total"):
        return 0
    duck = Duck(storage)
    try:
        duck.view("sc", f"{layout.STATEMENT_CHECKS}/*/*.parquet")
        have_companies = duck.view("companies", layout.COMPANIES, hive=False)
        name_col = "c.name" if have_companies else "NULL"
        name_join = "LEFT JOIN companies c ON c.cik = q.cik" if have_companies else ""
        rows = duck.fetch_dicts(
            _base_sql()
            + f"""
            SELECT q.cik, {name_col} AS name, q.accession, q.statement, q.check_name, q.reason,
                   q.lhs, q.rhs, q.lhs - q.rhs AS difference, round(q.q, 2) AS tolerance_multiple, q.detail
            FROM r q {name_join}
            WHERE NOT q.passed
            ORDER BY q.reason, q.check_name, q.cik, q.accession
            """
        )
    finally:
        duck.close()

    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="44546A")
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Failing arithmetic checks"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = (
        f"{report['failed']:,} of {report['total']:,} checks fail "
        f"({report['failed'] / report['total'] * 100:.1f} %). "
        f"{report['companies_with_failure']:,} of {report['companies_checked']:,} companies have at least one."
    )
    summary["A3"] = "Not every failure is ours: a filer's own tagging error is listed, never silently corrected."

    at = 5
    for title, data in (
        (
            "By check",
            [
                (CHECK_LABELS.get(r["check_name"], r["check_name"]), r["failed"], r["total"])
                for r in report["per_check"]
            ],
        ),
        ("By reason", None),
    ):
        summary.cell(at, 1, title).font = Font(bold=True)
        at += 1
        if data is None:
            counts: dict[str, int] = {}
            for row in rows:
                counts[row["reason"]] = counts.get(row["reason"], 0) + 1
            data = [(k, v, report["failed"]) for k, v in sorted(counts.items(), key=lambda kv: -kv[1])]
        for cell, value in zip("ABC", ("", "failed", "of which ran"), strict=True):
            summary[f"{cell}{at}"] = value
            summary[f"{cell}{at}"].font = head
            summary[f"{cell}{at}"].fill = fill
        at += 1
        for label, failed, ran in data:
            summary.cell(at, 1, label)
            summary.cell(at, 2, failed).number_format = "#,##0"
            summary.cell(at, 3, ran).number_format = "#,##0"
            at += 1
        at += 1
    summary.column_dimensions["A"].width = 54
    summary.column_dimensions["B"].width = 12
    summary.column_dimensions["C"].width = 14

    ws = wb.create_sheet("Failures")
    headers = [
        ("Company", 34),
        ("CIK", 10),
        ("Filing", 22),
        ("Statement", 10),
        ("Check", 46),
        ("Why", 42),
        ("Computed", 20),
        ("Reported", 20),
        ("Difference", 20),
        ("Times over tolerance", 20),
        ("What was compared", 80),
    ]
    for i, (title, width) in enumerate(headers, start=1):
        cell = ws.cell(1, i, title)
        cell.font, cell.fill, cell.alignment = head, fill, Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    for n, row in enumerate(rows, start=2):
        ws.cell(n, 1, row["name"] or "")
        ws.cell(n, 2, str(row["cik"]))  # text, so it is not reformatted as a number
        ws.cell(n, 3, row["accession"] or "")
        ws.cell(n, 4, row["statement"] or "")
        ws.cell(n, 5, CHECK_LABELS.get(row["check_name"], row["check_name"]))
        ws.cell(n, 6, row["reason"])
        for col, key in ((7, "lhs"), (8, "rhs"), (9, "difference")):
            ws.cell(n, col, row[key]).number_format = "#,##0.00"
        ws.cell(n, 10, row["tolerance_multiple"]).number_format = "#,##0.00"
        ws.cell(n, 11, row["detail"] or "")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    wb.save(path)
    return len(rows)
