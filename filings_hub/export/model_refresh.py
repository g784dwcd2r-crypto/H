"""Preview or apply guarded financial-data mappings to a new, supported .xlsx workbook.

This is a bounded local CLI prototype, not a native Excel add-in or arbitrary-workbook updater.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import math
import os
import re
import tempfile
import warnings
import zipfile
from datetime import date
from pathlib import Path
from typing import Any, Literal
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.workbook.properties import CalcProperties
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator

from filings_hub.db.database import Database
from filings_hub.platform.financials import financial_snapshot

SCALES = {"units": 1, "thousands": 1000, "millions": 1_000_000, "billions": 1_000_000_000}
PARTS = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|"
    r"xl/(?:workbook|styles|sharedStrings|calcChain)\.xml|xl/_rels/workbook\.xml\.rels|"
    r"xl/theme/theme\d+\.xml|xl/worksheets/sheet\d+\.xml|"
    r"xl/worksheets/_rels/sheet\d+\.xml\.rels|xl/comments/comment\d+\.xml|"
    r"xl/comments\d+\.xml|xl/drawings/commentsDrawing\d+\.vml)"
)
UNSUPPORTED_XML = {"extLst", "AlternateContent", "oleObjects", "controls", "drawing", "picture", "legacyDrawingHF"}


class ModelRefreshError(ValueError):
    pass


class UnsupportedWorkbook(ModelRefreshError):
    pass


class CellMapping(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    sheet: str = Field(min_length=1, max_length=31)
    cell: str
    cik: int = Field(gt=0, lt=10**10)
    statement: Literal["IS", "BS", "CF", "EQ", "CI"]
    line_key: str = Field(min_length=1, max_length=1000)
    period_label: str = Field(min_length=1, max_length=50)
    unit: str = Field(min_length=1, max_length=100)
    scale: Literal["units", "thousands", "millions", "billions"] = "units"
    last_written: float | int | None

    @field_validator("cell")
    @classmethod
    def valid_cell(cls, value: str) -> str:
        if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", value):
            raise ValueError("cell must be one absolute worksheet coordinate such as B4 (without $)")
        row, column = coordinate_to_tuple(value)
        if row > 1048576 or column > 16384:
            raise ValueError("cell is outside Excel's supported worksheet bounds")
        return value

    @field_validator("last_written", mode="before")
    @classmethod
    def numeric_guard(cls, value):
        if value is not None and (type(value) not in (int, float) or not math.isfinite(value)):
            raise ValueError("last_written must be a finite number or explicit null for an empty target")
        return value

    @model_validator(mode="after")
    def scale_unit(self):
        if self.scale != "units" and not re.fullmatch("[A-Z]{3}", self.unit):
            raise ValueError("only currency totals may be scaled; shares/per-share/ratios must use units")
        return self


class ModelMapping(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    schema_version: Literal[1] = 1
    period_mode: Literal["as_filed", "annual", "quarterly", "ltm"] = "annual"
    presentation: Literal["original", "latest"] = "original"
    as_of: str | None = None
    mappings: list[CellMapping] = Field(min_length=1, max_length=500)

    @field_validator("as_of")
    @classmethod
    def valid_date(cls, value):
        if value is not None:
            date.fromisoformat(value)
        return value

    @model_validator(mode="after")
    def unique_targets(self):
        targets = [(mapping.sheet, mapping.cell) for mapping in self.mappings]
        if len(targets) != len(set(targets)):
            raise ValueError("each worksheet cell may have only one mapping")
        if self.presentation == "latest" and self.period_mode in ("quarterly", "ltm"):
            raise ValueError("latest presentation is supported only for annual/as_filed periods")
        return self


def inspect_package(raw: bytes) -> None:
    """Fail before openpyxl can silently discard unsupported package parts or extensions."""
    if len(raw) > 30 * 1024 * 1024:
        raise UnsupportedWorkbook("Workbook exceeds the 30 MiB input limit.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            items = package.infolist()
            if len(items) > 1000 or sum(item.file_size for item in items) > 200 * 1024 * 1024:
                raise UnsupportedWorkbook("Workbook exceeds the supported expanded package size or entry count.")
            if len({item.filename for item in items}) != len(items):
                raise UnsupportedWorkbook("Workbook contains duplicate package entries.")
            for item in items:
                if not PARTS.fullmatch(item.filename):
                    raise UnsupportedWorkbook(f"Unsupported workbook component: {item.filename}")
                if item.filename.endswith((".xml", ".vml", ".rels")):
                    element = ET.fromstring(package.read(item))
                    for node in element.iter():
                        tag = node.tag.rsplit("}", 1)[-1]
                        if tag in UNSUPPORTED_XML:
                            raise UnsupportedWorkbook(f"Unsupported workbook feature: {tag}")
                        if tag == "f" and node.get("t") in ("array", "dataTable"):
                            raise UnsupportedWorkbook("Array/spill and data-table formulas are not supported.")
                        if tag == "ClientData" and node.get("ObjectType") != "Note":
                            raise UnsupportedWorkbook("Only ordinary comments are supported in VML drawings.")
                        if (
                            item.filename.endswith(".rels")
                            and tag == "Relationship"
                            and node.get("TargetMode") == "External"
                            and not (node.get("Type") or "").endswith("/hyperlink")
                        ):
                            raise UnsupportedWorkbook("External workbook links/connections are not supported.")
    except (zipfile.BadZipFile, ET.ParseError, RuntimeError) as exc:
        raise UnsupportedWorkbook("Workbook is not a readable, supported OOXML package.") from exc


def _numeric(value) -> bool:
    return type(value) in (int, float) and math.isfinite(value)


def _same_guard(actual, expected) -> bool:
    # Int/float equivalence is fine; booleans and text never impersonate numbers. No tolerance hides manual edits.
    return actual is None if expected is None else _numeric(actual) and actual == expected


def _load_supported(raw: bytes):
    inspect_package(raw)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        workbook = load_workbook(io.BytesIO(raw), data_only=False, keep_links=True, rich_text=True)
    if caught:
        raise UnsupportedWorkbook("The workbook reader reported an unsupported feature: " + str(caught[0].message))
    if workbook.security.lockStructure or workbook.security.lockWindows:
        raise UnsupportedWorkbook("Protected workbook structure is not supported for model refresh.")
    if any(sheet.protection.sheet for sheet in workbook.worksheets):
        raise UnsupportedWorkbook("Protected worksheets are not supported for model refresh.")
    return workbook


def _cell_state(workbook) -> dict[tuple[str, str], Any]:
    return {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook
        for row in sheet
        for cell in row
        if not isinstance(cell, MergedCell) and cell.value is not None
    }


def refresh_model(
    database: Database,
    input_path: str | Path,
    mapping: ModelMapping | dict[str, Any],
    *,
    output_path: str | Path | None = None,
    apply: bool = False,
) -> dict[str, Any]:
    """All-or-none guarded update. Preview by default; even --apply creates no workbook with conflicts."""
    mapping = mapping if isinstance(mapping, ModelMapping) else ModelMapping.model_validate(mapping)
    source = Path(input_path).resolve()
    if source.suffix.lower() != ".xlsx":
        raise UnsupportedWorkbook("Only ordinary .xlsx workbooks are supported; no macro-enabled/binary workbooks.")
    output = Path(output_path).resolve() if output_path is not None else None
    if apply and output is None:
        raise ModelRefreshError("Applying a refresh requires a new output workbook path.")
    if output is not None:
        if output == source or output.exists():
            raise ModelRefreshError("Output must be a new file; the input and existing files are never replaced.")
        if output.suffix.lower() != ".xlsx":
            raise ModelRefreshError("Output must use the .xlsx extension.")
    raw = source.read_bytes()
    workbook = _load_supported(raw)
    original_state = _cell_state(workbook)
    original_hash = hashlib.sha256(raw).hexdigest()
    snapshots = {}
    with database.read_snapshot() as reader:
        for cik in sorted({item.cik for item in mapping.mappings}):
            periods = sorted({item.period_label for item in mapping.mappings if item.cik == cik})
            snapshots[cik] = financial_snapshot(
                reader,
                cik,
                periods=periods,
                limit=len(periods),
                period_mode=mapping.period_mode,
                presentation=mapping.presentation,
                as_of=date.fromisoformat(mapping.as_of) if mapping.as_of else None,
            )
    changes, prepared = [], []
    for item in mapping.mappings:
        snapshot = snapshots[item.cik]
        entry = {
            **item.model_dump(),
            "snapshot_id": snapshot["snapshot_id"],
            "current_value": None,
            "proposed_value": None,
            "status": "conflict",
            "reason": None,
            "evidence": None,
        }
        changes.append(entry)
        if item.sheet not in workbook.sheetnames:
            entry["reason"] = "Mapped worksheet does not exist."
            continue
        sheet = workbook[item.sheet]
        cell = sheet[item.cell]
        entry["current_value"] = (
            cell.value if cell.value is None or type(cell.value) in (str, int, float, bool) else str(cell.value)
        )
        if isinstance(cell, MergedCell) or any(item.cell in area for area in sheet.merged_cells.ranges):
            entry["reason"] = "Merged-cell targets are not supported."
            continue
        if cell.data_type == "f":
            entry["reason"] = "Target contains a formula; formulas are never overwritten."
            continue
        if not _same_guard(cell.value, item.last_written):
            entry["reason"] = "Cell differs from last_written; preserve the manual edit and review its mapping."
            continue
        statement = next((s for s in snapshot["grid"]["statements"] if s["code"] == item.statement), None)
        lines = [line for line in statement["lines"] if line["key"] == item.line_key] if statement else []
        if len(lines) != 1 or lines[0]["is_abstract"]:
            entry["reason"] = "Mapped financial line is missing or ambiguous; no guessed replacement was used."
            continue
        line = lines[0]
        if line["unit"] != item.unit:
            entry["reason"] = f"Source unit changed or differs: expected {item.unit}, got {line['unit']}."
            continue
        value = line["values"].get(item.period_label)
        evidence = line["value_metadata"].get(item.period_label)
        if (
            not _numeric(value)
            or not evidence
            or not evidence.get("sources")
            or evidence.get("status") == "unavailable"
        ):
            entry["status"] = "unavailable"
            entry["reason"] = "A finite value with source evidence is unavailable; the workbook is unchanged."
            continue
        scaled = value / SCALES[item.scale]
        entry.update(proposed_value=scaled, status="unchanged" if cell.value == scaled else "change", evidence=evidence)
        prepared.append((item, cell, entry))
    conflicts = sum(entry["status"] in ("conflict", "unavailable") for entry in changes)
    report = dict(
        status="blocked" if conflicts else "ready",
        applied=False,
        input_sha256=original_hash,
        input_path=str(source),
        output_path=str(output) if output else None,
        conflicts=conflicts,
        changes=changes,
        next_mapping=None,
        source_snapshots={str(cik): snap["snapshot_id"] for cik, snap in snapshots.items()},
        formula_recalculation_required=True,
        limitations=[
            "This is an explicit-cell local refresh prototype, not a native Excel add-in.",
            "Formula expressions are preserved; cached results are not calculated by openpyxl.",
            "Recalculate and review the output in Excel before relying on formula results.",
            *next(iter(snapshots.values()))["limitations"],
        ],
    )
    if not apply or conflicts:
        return report
    next_mapping = copy.deepcopy(mapping.model_dump())
    for item, cell, entry in prepared:
        cell.value = entry["proposed_value"]
        evidence_text = "Disclosure source evidence\n" + json.dumps(
            {
                "snapshot_id": entry["snapshot_id"],
                "cik": item.cik,
                "statement": item.statement,
                "line_key": item.line_key,
                "period_label": item.period_label,
                "unit": item.unit,
                "scale": item.scale,
                "evidence": entry["evidence"],
            },
            ensure_ascii=False,
            sort_keys=True,
        )
        old_comment = cell.comment
        if old_comment and old_comment.author != "Disclosure":
            evidence_text = old_comment.text + "\n\n" + evidence_text
        cell.comment = Comment(
            evidence_text, old_comment.author if old_comment and old_comment.author != "Disclosure" else "Disclosure"
        )
        for next_item in next_mapping["mappings"]:
            if (next_item["sheet"], next_item["cell"]) == (item.sheet, item.cell):
                next_item["last_written"] = cell.value
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    buffer = io.BytesIO()
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        workbook.save(buffer)
    if caught:
        raise UnsupportedWorkbook("Workbook save would lose an unsupported feature: " + str(caught[0].message))
    # Verify formulas/manual cells and each mapped value after real serialization, before publishing output.
    saved = load_workbook(io.BytesIO(buffer.getvalue()), data_only=False, rich_text=True)
    expected_state = dict(original_state)
    for item, _cell, entry in prepared:
        expected_state[item.sheet, item.cell] = entry["proposed_value"]
        if (
            not saved[item.sheet][item.cell].comment
            or entry["snapshot_id"] not in saved[item.sheet][item.cell].comment.text
        ):
            raise ModelRefreshError("Source evidence did not survive workbook serialization.")
    if _cell_state(saved) != expected_state:
        raise ModelRefreshError(
            "Workbook serialization changed a formula, manual value or mapped value; no output written."
        )
    if hashlib.sha256(source.read_bytes()).hexdigest() != original_hash:
        raise ModelRefreshError("The input changed during refresh; no stale output was written.")
    assert output is not None
    output.parent.mkdir(parents=True, exist_ok=True)
    # Link a complete temporary artifact into a previously absent name. No partial workbook is published,
    # and a concurrent run cannot replace an existing output or symlink.
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output.parent, prefix=".disclosure-refresh-", suffix=".xlsx", delete=False
        ) as handle:
            temporary = Path(handle.name)
            handle.write(buffer.getvalue())
            handle.flush()
            os.fsync(handle.fileno())
        os.link(temporary, output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    report.update(
        status="written",
        applied=True,
        next_mapping=next_mapping,
        output_sha256=hashlib.sha256(buffer.getvalue()).hexdigest(),
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("mapping", type=Path)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--apply", action="store_true", help="apply only if every target and guard passes")
    args = parser.parse_args()
    from filings_hub.config import get_settings
    from filings_hub.lake.storage import Storage

    settings = get_settings()
    database = Database(settings.database_url, Storage(settings.resolved_lake_root()))
    try:
        report = refresh_model(
            database, args.input, json.loads(args.mapping.read_text()), output_path=args.out, apply=args.apply
        )
        print(json.dumps(report, indent=2, ensure_ascii=False, allow_nan=False))
        if report["conflicts"]:
            raise SystemExit(2)
    finally:
        database.close()


if __name__ == "__main__":
    main()
