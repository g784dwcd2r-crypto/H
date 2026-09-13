from __future__ import annotations

import copy
import io
import json
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Font
from pydantic import ValidationError

from filings_hub.export.model_refresh import ModelMapping, ModelRefreshError, UnsupportedWorkbook, refresh_model
from filings_hub.platform.financials import financial_snapshot
from filings_hub.testing import edgar_fixtures as fx


@pytest.fixture
def model(tmp_path, db):
    source = tmp_path / "analyst-model.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Model"
    sheet["A2"], sheet["B2"] = "Revenue (millions)", 300000
    sheet["A3"], sheet["B3"] = "EPS", 7
    sheet["B4"] = "=B2/B3"
    sheet["C1"] = 123.45
    sheet["C1"].comment = Comment("Analyst assumption: hold constant", "Analyst")
    sheet["B2"].font = Font(bold=True, color="123456")
    sheet["B2"].number_format = "#,##0.00"
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 25
    workbook.save(source)
    snapshot = financial_snapshot(db, fx.APPLE, periods=["FY2025"], period_mode="annual")
    lines = next(statement["lines"] for statement in snapshot["grid"]["statements"] if statement["code"] == "IS")
    revenue = next(line for line in lines if line["concept"] == "RevenueFromContractWithCustomerExcludingAssessedTax")
    eps = next(line for line in lines if line["concept"] == "EarningsPerShareDiluted")
    mapping = {
        "schema_version": 1,
        "period_mode": "annual",
        "presentation": "original",
        "as_of": None,
        "mappings": [
            {
                "sheet": "Model",
                "cell": "B2",
                "cik": fx.APPLE,
                "statement": "IS",
                "line_key": revenue["key"],
                "period_label": "FY2025",
                "unit": "USD",
                "scale": "millions",
                "last_written": 300000,
            },
            {
                "sheet": "Model",
                "cell": "B3",
                "cik": fx.APPLE,
                "statement": "IS",
                "line_key": eps["key"],
                "period_label": "FY2025",
                "unit": "USD/shares",
                "scale": "units",
                "last_written": 7,
            },
        ],
    }
    return source, mapping


def test_preview_is_concrete_and_never_writes(db, model, tmp_path):
    source, mapping = model
    before = source.read_bytes()
    output = tmp_path / "new.xlsx"
    report = refresh_model(db, source, mapping, output_path=output)
    assert report["status"] == "ready" and not report["applied"]
    assert [item["proposed_value"] for item in report["changes"]] == [416161, 7.46]
    assert all(item["evidence"]["sources"] for item in report["changes"])
    assert not output.exists() and source.read_bytes() == before
    assert report["next_mapping"] is None
    json.dumps(report, allow_nan=False)


def test_update_roundtrip_preserves_formulas_manual_values_styles_and_source_evidence(db, model, tmp_path):
    source, mapping = model
    before = source.read_bytes()
    output = tmp_path / "refreshed.xlsx"
    report = refresh_model(db, source, mapping, output_path=output, apply=True)
    assert report["status"] == "written" and report["applied"] and report["conflicts"] == 0
    assert source.read_bytes() == before
    workbook = load_workbook(output, data_only=False)
    sheet = workbook["Model"]
    assert sheet["B2"].value == 416161 and sheet["B3"].value == 7.46
    assert sheet["B4"].value == "=B2/B3" and sheet["C1"].value == 123.45
    assert sheet["C1"].comment.text == "Analyst assumption: hold constant"
    assert sheet["C1"].comment.author == "Analyst"
    assert sheet["B2"].font.bold and sheet["B2"].font.color.rgb == "00123456"
    assert sheet["B2"].number_format == "#,##0.00" and sheet.freeze_panes == "B2"
    assert sheet.column_dimensions["A"].width == 25
    assert fx.APPLE_10K_FY2025 in sheet["B2"].comment.text
    assert report["source_snapshots"][str(fx.APPLE)] in sheet["B2"].comment.text
    assert workbook.calculation.fullCalcOnLoad and report["formula_recalculation_required"]
    assert [item["last_written"] for item in report["next_mapping"]["mappings"]] == [416161, 7.46]
    # An ordinary second pass using guarded next_mapping is valid and preserves the first evidence.
    again = refresh_model(db, output, report["next_mapping"], output_path=tmp_path / "again.xlsx", apply=True)
    assert again["applied"] and all(item["status"] == "unchanged" for item in again["changes"])


@pytest.mark.parametrize("manual", [999, True, "300000", "=1+2"])
def test_manual_overrides_and_target_formulas_block_every_update(db, model, tmp_path, manual):
    source, mapping = model
    workbook = load_workbook(source)
    workbook["Model"]["B2"] = manual
    workbook.save(source)
    before = source.read_bytes()
    output = tmp_path / "blocked.xlsx"
    report = refresh_model(db, source, mapping, output_path=output, apply=True)
    assert report["status"] == "blocked" and not report["applied"] and report["conflicts"] == 1
    assert report["changes"][0]["status"] == "conflict"
    assert report["changes"][1]["status"] == "change"  # valid proposals are still reviewable, none applied
    assert not output.exists() and source.read_bytes() == before


@pytest.mark.parametrize(
    "field,value", [("unit", "EUR"), ("line_key", "nonexistent"), ("period_label", "FY1900"), ("sheet", "Missing")]
)
def test_missing_source_changed_unit_and_missing_targets_are_explicit(db, model, tmp_path, field, value):
    source, mapping = model
    mapping["mappings"][0][field] = value
    report = refresh_model(db, source, mapping, output_path=tmp_path / "blocked.xlsx", apply=True)
    assert report["status"] == "blocked" and report["changes"][0]["reason"]
    assert not (tmp_path / "blocked.xlsx").exists()


def test_cutoff_does_not_pull_later_filed_period_into_model(db, model):
    source, mapping = model
    mapping["as_of"] = "2024-01-01"
    report = refresh_model(db, source, mapping)
    assert report["status"] == "blocked" and report["conflicts"] == 2
    assert all(item["proposed_value"] is None for item in report["changes"])


def test_empty_guard_can_fill_empty_target_and_merged_target_is_rejected(db, model, tmp_path):
    source, mapping = model
    mapping["mappings"][0]["last_written"] = None
    workbook = load_workbook(source)
    workbook["Model"]["B2"] = None
    workbook.save(source)
    assert refresh_model(db, source, mapping)["status"] == "ready"
    workbook["Model"].merge_cells("B2:C2")
    workbook.save(source)
    report = refresh_model(db, source, mapping)
    assert report["status"] == "blocked" and "Merged" in report["changes"][0]["reason"]


@pytest.mark.parametrize("feature", ["chart", "protection", "external-part"])
def test_unsupported_workbook_objects_fail_before_any_output(db, model, tmp_path, feature):
    source, mapping = model
    workbook = load_workbook(source)
    if feature == "chart":
        chart = BarChart()
        chart.add_data(Reference(workbook["Model"], min_col=2, min_row=2, max_row=3))
        workbook["Model"].add_chart(chart, "F1")
        workbook.save(source)
    elif feature == "protection":
        workbook["Model"].protection.sheet = True
        workbook.save(source)
    else:
        with zipfile.ZipFile(source, "a") as package:
            package.writestr("xl/vbaProject.bin", b"unsupported macro content")
    before = source.read_bytes()
    with pytest.raises(UnsupportedWorkbook):
        refresh_model(db, source, mapping, output_path=tmp_path / "unsupported.xlsx", apply=True)
    assert source.read_bytes() == before and not (tmp_path / "unsupported.xlsx").exists()


def test_array_formula_package_is_rejected(db, model):
    source, mapping = model
    buffer = io.BytesIO()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(buffer, "w") as changed:
        for part in original.infolist():
            raw = original.read(part)
            if part.filename == "xl/worksheets/sheet1.xml":
                raw = raw.replace(b"<f>", b'<f t="array" ref="B4">')
            changed.writestr(part, raw)
    source.write_bytes(buffer.getvalue())
    with pytest.raises(UnsupportedWorkbook, match="Array/spill"):
        refresh_model(db, source, mapping)


def test_input_and_existing_outputs_cannot_be_replaced(db, model, tmp_path):
    source, mapping = model
    with pytest.raises(ModelRefreshError, match="never replaced"):
        refresh_model(db, source, mapping, output_path=source, apply=True)
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"valuable existing content")
    with pytest.raises(ModelRefreshError, match="never replaced"):
        refresh_model(db, source, mapping, output_path=output, apply=True)
    assert output.read_bytes() == b"valuable existing content"
    with pytest.raises(ModelRefreshError, match="requires a new output"):
        refresh_model(db, source, mapping, apply=True)


@pytest.mark.parametrize(
    "mutation", ["duplicate", "scale_eps", "missing_guard", "bad_coordinate", "nan", "unknown_field"]
)
def test_mapping_validation_rejects_ambiguous_or_unsafe_configuration(model, mutation):
    _, original = model
    mapping = copy.deepcopy(original)
    if mutation == "duplicate":
        mapping["mappings"].append(mapping["mappings"][0])
    elif mutation == "scale_eps":
        mapping["mappings"][1]["scale"] = "millions"
    elif mutation == "missing_guard":
        del mapping["mappings"][0]["last_written"]
    elif mutation == "bad_coordinate":
        mapping["mappings"][0]["cell"] = "XFE1"
    elif mutation == "nan":
        mapping["mappings"][0]["last_written"] = float("nan")
    elif mutation == "unknown_field":
        mapping["mappings"][0]["overwrite_formulas"] = True
    with pytest.raises(ValidationError):
        ModelMapping.model_validate(mapping)
