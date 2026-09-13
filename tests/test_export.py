import io

import pytest
from openpyxl import load_workbook

from filings_hub.db.database import Database
from filings_hub.export.excel import ExportOptions, export_excel, export_workbook, workbook_from_grid
from filings_hub.export.grid import Grid, build_grid, merge_line_order
from filings_hub.testing import edgar_fixtures as fx


def test_merge_line_order():
    assert merge_line_order([["a", "b", "c"], ["a", "x", "c"]]) == ["a", "x", "b", "c"]
    assert merge_line_order([["a", "b"], ["z", "a"]]) == ["z", "a", "b"]
    assert merge_line_order([[], ["q"]]) == ["q"]
    assert merge_line_order([["a", "b"], ["b", "a"]]) == ["a", "b"]  # newest order wins


def test_grid_apple(db: Database):
    g = build_grid(db, fx.APPLE, ["FY2025", "Q1 2026", "Q2 2026"])
    assert [p.period_label for p in g.periods] == ["FY2025", "Q1 2026", "Q2 2026"]
    assert g.periods[2].is_provisional and not g.periods[1].is_provisional
    codes = {s.code: s for s in g.statements}
    assert set(codes) == {"IS", "BS", "CF"}
    rev = next(ln for ln in codes["IS"].lines if ln.concept == "RevenueFromContractWithCustomerExcludingAssessedTax")
    assert rev.values == {
        "FY2025": 416_161_000_000.0,
        "Q1 2026": 140_000_000_000.0,
        "Q2 2026": 100_000_000_000.0,
    }
    assert rev.label == "Total net sales" and rev.is_subtotal
    d = g.to_dict()
    assert (
        d["periods"][0]["accession"] == fx.APPLE_10K_FY2025
        and d["statements"][0]["lines"][1]["values"]["FY2025"] == 416_161_000_000.0
    )


def test_grid_default_limit_and_unknown(db: Database):
    g = build_grid(db, fx.APPLE, None, limit=3)
    assert len(g.periods) == 3 and g.periods[-1].period_label == "Q3 2026"
    with pytest.raises(KeyError):
        build_grid(db, 999999)


def test_excel_roundtrip_matches_statements(db: Database):
    data = export_excel(db, fx.APPLE, ["FY2024", "FY2025", "Q1 2026"])
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Income Statement", "Balance Sheet", "Cash Flow", "Source"]
    ws = wb["Income Statement"]
    assert ws["A1"].value == "Apple Inc. (AAPL)"
    assert ws.column_dimensions["B"].hidden
    header = [c.value for c in ws[4]]
    assert header[:3] == ["Line item", "Concept", "Unit"] and header[3:] == [
        "FY2024 (provisional)",
        "FY2025",
        "Q1 2026",
    ]
    rows = {r[1].value: r for r in ws.iter_rows(min_row=7) if r[1].value}
    ni = rows["NetIncomeLoss"]
    assert [c.value for c in ni[3:]] == [93_736_000_000, 112_861_000_000, 41_000_000_000]
    assert all(isinstance(c.value, int | float) for c in ni[3:])  # numbers, not text
    assert ni[3].number_format.startswith("#,##0")
    eps = rows["EarningsPerShareDiluted"]
    assert eps[2].value == "USD/shares" and eps[4].value == 7.46 and eps[4].number_format.startswith("0.00")
    cf = wb["Cash Flow"]
    cf_rows = {r[1].value: r for r in cf.iter_rows(min_row=7) if r[1].value}
    assert cf_rows["PaymentsForRepurchaseOfCommonStock"][4].value == -95_000_000_000  # negatives as negatives
    # column totals: net income equals what the statements table holds for each filing
    for j, acc in enumerate([fx.APPLE_10K_FY2024, fx.APPLE_10K_FY2025, fx.APPLE_10Q_Q1_2026]):
        v = db.query(
            "SELECT value_presented FROM statements WHERE accession = ? AND statement = 'IS' AND concept = 'NetIncomeLoss' AND is_primary_period",
            [acc],
        )
        assert ni[3 + j].value == v[0]["value_presented"]
    src = wb["Source"]
    assert src["A2"].value == "FY2024" and src["E3"].value == fx.APPLE_10K_FY2025 and src["C2"].value == "as filed"
    assert src["G3"].hyperlink.target.endswith("-index.htm")
    assert src["J2"].value.startswith("provisional") and src["K3"].value == "passed"
    # no formulas anywhere
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith("="))


def test_workbook_with_no_statements():
    g = Grid(cik=1, company_name="Empty Co", ticker=None, periods=[], statements=[])
    wb = workbook_from_grid(g)
    assert wb.sheetnames == ["No statements", "Source"]


def test_repeated_concepts_key_on_their_label(db: Database):
    """A concept presented twice -- "beginning balances" and "ending balances" carry the same concept --
    must key on what distinguishes the lines. Keying by position alone meant that when one filing
    presented fewer occurrences than another, every later occurrence shifted up and the values landed on
    the wrong row."""
    from filings_hub.export.grid import _keyed_lines

    rows = [
        {"concept": "Cash", "label": "Cash, beginning balances"},
        {"concept": "Ops", "label": "Operating activities"},
        {"concept": "Cash", "label": "Cash, ending balances"},
    ]
    keys = [k for k, _ in _keyed_lines(rows)]
    assert len(set(keys)) == 3
    # dropping the opening line must not rename the closing one
    without_opening = [k for k, _ in _keyed_lines(rows[1:])]
    assert without_opening[-1] == keys[-1]
    # a genuinely identical line (same concept and label) still gets a distinct key
    twice = [k for k, _ in _keyed_lines([rows[0], rows[0]])]
    assert twice[0] != twice[1]


def test_grid_aligns_opening_and_closing_cash_across_filings(db: Database):
    """The cash flow statement's opening and closing balances share one concept: each must keep its own
    row across every period column."""
    g = build_grid(db, fx.APPLE, ["FY2025", "Q1 2026", "Q2 2026"])
    cf = next(s for s in g.statements if s.code == "CF")
    cash = [ln for ln in cf.lines if ln.concept == fx.CASH_CONCEPT]
    assert len(cash) == 2
    opening, closing = cash
    assert "beginning" in opening.label and "ending" in closing.label
    for label, key in (("FY2025", fx.P_FY2025), ("Q1 2026", fx.P_Q1_2026), ("Q2 2026", fx.P_H1_2026)):
        want_open, want_close = fx.APPLE_CASH[key]
        assert opening.values[label] == float(want_open), (label, "opening")
        assert closing.values[label] == float(want_close), (label, "closing")
    # and each column reconciles against its own change line
    change = next(ln for ln in cf.lines if ln.label.startswith("Increase/(Decrease)"))
    for label in ("FY2025", "Q1 2026", "Q2 2026"):
        assert opening.values[label] + change.values[label] == closing.values[label], label


def _line(g: Grid, code: str, concept: str):
    stmt = next(s for s in g.statements if s.code == code)
    return next(ln for ln in stmt.lines if ln.concept == concept)


def test_grid_restated_takes_the_latest_comparative(db: Database):
    as_filed = build_grid(db, fx.APPLE, ["FY2024", "FY2025", "Q1 2026"])
    restated = build_grid(db, fx.APPLE, ["FY2024", "FY2025", "Q1 2026"], restated=True)
    rev = "RevenueFromContractWithCustomerExcludingAssessedTax"
    assert _line(as_filed, "IS", rev).values["FY2024"] == 391_035_000_000.0  # the provisional 10-K's own number
    assert _line(restated, "IS", rev).values["FY2024"] == 391_036_000_000.0  # as presented in the FY2025 10-K
    col = next(p for p in restated.periods if p.period_label == "FY2024")
    assert col.basis == "restated" and col.restated_from == fx.APPLE_10K_FY2025
    assert restated.periods[-1].basis == "as filed"  # nothing later presents Q1 2026
    d = restated.to_dict()
    assert d["restated"] is True and d["periods"][0]["basis_note"].endswith(fx.APPLE_10K_FY2025)


def test_grid_annual_and_column_order(db: Database):
    g = build_grid(db, fx.APPLE, None, limit=5, period_mode="annual")
    assert [p.period_label for p in g.periods] == ["FY2024", "FY2025"]
    left = build_grid(db, fx.APPLE, None, limit=3, column_order="newest_left")
    assert [p.period_label for p in left.periods] == ["Q3 2026", "Q2 2026", "Q1 2026"]
    assert left.column_order == "newest_left"
    with pytest.raises(ValueError):
        build_grid(db, fx.APPLE, period_mode="weekly")


def test_grid_quarterly_derives_cash_flow_quarters(db: Database):
    g = build_grid(db, fx.APPLE, None, limit=8, period_mode="quarterly")
    labels = [p.period_label for p in g.periods]
    assert "Q4 2025" in labels and "FY2025" not in labels
    q4 = next(p for p in g.periods if p.period_label == "Q4 2025")
    assert q4.basis == "derived" and q4.filed_label == "FY2025"
    ocf = _line(g, "CF", "NetCashProvidedByUsedInOperatingActivities")
    # 10-Q cash flow statements are year to date: Q2 = six months less three months
    assert ocf.values["Q1 2025"] == 29_935_000_000.0
    assert ocf.values["Q2 2025"] == pytest.approx(53_912_000_000.0 - 29_935_000_000.0)
    # the income statement already presents the quarter: unchanged
    assert _line(g, "IS", "NetIncomeLoss").values["Q2 2025"] == 24_780_000_000.0
    # the opening cash balance of a derived quarter is the prior quarter's closing balance
    cash = [ln for ln in next(s for s in g.statements if s.code == "CF").lines if ln.concept.startswith("CashCash")]
    assert cash[0].values["Q2 2026"] == cash[-1].values["Q1 2026"] == 32_000_000_000.0
    # the fixture has no Q3 statements, so Q4 cannot be derived and stays empty rather than wrong
    assert ocf.values["Q4 2025"] is None and q4.basis_note
    # balance sheets are points in time and stay as filed under the Q4 label
    assert _line(g, "BS", "Assets").values["Q4 2025"] == 360_000_000_000.0


def test_grid_ltm(db: Database):
    g = build_grid(db, fx.APPLE, None, limit=8, period_mode="ltm")
    ni = _line(g, "IS", "NetIncomeLoss")
    # YTD Q2 2026 + FY2025 - YTD Q2 2025 (the comparative in the Q2 2026 10-Q)
    assert ni.values["LTM Q2 2026"] == pytest.approx(65_100_000_000.0 + 112_861_000_000.0 - 61_110_000_000.0)
    assert ni.values["FY2025"] == 112_861_000_000.0
    col = next(p for p in g.periods if p.period_label == "LTM Q2 2026")
    assert col.basis == "derived" and col.filed_label == "Q2 2026"
    assert _line(g, "BS", "Assets").values["LTM Q2 2026"] == 372_000_000_000.0


def test_export_options_validation():
    assert ExportOptions.from_dict(None) == ExportOptions()
    o = ExportOptions.from_dict(
        {"layout": "one_sheet", "include_source": "false", "statements": "IS,BS", "scale": "millions"}
    )
    assert o.layout == "one_sheet" and o.include_source is False and o.statements == ("IS", "BS")
    for bad in ({"layout": "pdf"}, {"scale": "lakhs"}, {"statements": "XX"}, {"filename": "{owner}"}):
        with pytest.raises(ValueError):
            ExportOptions.from_dict(bad)


def test_export_one_sheet_periods_down_formulas_and_scale(db: Database):
    opts = ExportOptions(
        layout="one_sheet",
        orientation="periods_down",
        subtotals="formulas",
        scale="millions",
        negative_style="minus",
        include_concepts=False,
        filename="{ticker}-{mode}-{periods}p",
        statements=("IS", "CF"),
    )
    data, fname = export_workbook(db, fx.APPLE, ["FY2025", "Q1 2026"], options=opts)
    assert fname == "AAPL-as_filed-2p.xlsx"
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Statements", "Source"]
    ws = wb["Statements"]
    assert ws["A1"].value == "Apple Inc. (AAPL) - Income Statement"
    header = [c.value for c in ws[4]]
    assert header[:3] == ["Period", "Period end", "Filing"] and "Total net sales" in header
    # one row per period; revenue in millions
    rev_col = header.index("Total net sales") + 1
    assert ws.cell(5, 1).value == "FY2025" and ws.cell(5, rev_col).value == 416_161.0
    assert ws.cell(5, rev_col).number_format.startswith("#,##0.0;-")
    # gross profit is a subtotal: written as a formula only when its children add up to the reported number
    cells = [c.value for row in ws.iter_rows(min_row=5, max_row=6) for c in row]
    formulas = [v for v in cells if isinstance(v, str) and v.startswith("=")]
    assert all(v.startswith("=") and "+" in v for v in formulas)
    src = wb["Source"]
    assert src["C2"].value == "as filed"


def test_export_formulas_match_values(db: Database):
    plain = build_grid(db, fx.APPLE, ["FY2025"])
    wb = workbook_from_grid(plain, ExportOptions(subtotals="formulas"))
    ws = wb["Income Statement"]
    stmt = next(s for s in plain.statements if s.code == "IS")
    rows = {r[1].value: r for r in ws.iter_rows(min_row=7) if r[1].value}
    # every subtotal whose children add up carries a formula summing exactly those children
    for ln in stmt.lines:
        if not ln.is_subtotal:
            continue
        kids = [k for k in stmt.lines if k.parent_concept == ln.concept and not k.is_abstract]
        v = rows[ln.concept][3].value
        total, parts = ln.values["FY2025"], [k.values["FY2025"] for k in kids]
        if kids and total is not None and None not in parts and abs(sum(parts) - total) <= 1:
            assert isinstance(v, str) and v.count("+") == len(kids) - 1
        else:
            assert v == total


def test_export_quarterly_and_restated_workbook(db: Database):
    data, _ = export_workbook(db, fx.APPLE, None, 8, period_mode="quarterly", column_order="newest_left")
    wb = load_workbook(io.BytesIO(data))
    header = [c.value for c in wb["Cash Flow"][4]]
    assert header[3] == "Q3 2026" and "Q4 2025" in header  # newest left; Q4 derived from the fiscal year
    src = wb["Source"]
    basis = {src.cell(i, 1).value: src.cell(i, 3).value for i in range(2, 10)}
    assert basis["Q4 2025"].startswith("derived")
    data, _ = export_workbook(
        db, fx.APPLE, ["FY2024", "FY2025"], restated=True, options=ExportOptions(include_checks=False)
    )
    wb = load_workbook(io.BytesIO(data))
    assert wb["Income Statement"]["D4"].value == "FY2024 (provisional) (restated)"
    assert [c.value for c in wb["Source"][1]][-1] == "Statements source"
